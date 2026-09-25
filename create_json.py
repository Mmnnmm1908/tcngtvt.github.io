import os
import openpyxl
import json
import shutil
from PIL import Image
import io

excel_file = 'qr.20.08.xlsx'
output_json = 'data.json'
output_dir = 'anh-qr'

if not os.path.exists(excel_file):
    print(f"Loi: Khong tim thay file '{excel_file}' trong thu muc!")
else:
    # 1. Tự động dọn dẹp và tạo mới thư mục ảnh sạch sẽ
    if os.path.exists(output_dir):
        shutil.rmtree(output_dir)
    os.makedirs(output_dir)

    wb = openpyxl.load_workbook(excel_file, data_only=True)
    sheet = wb.active
    
    # 2. Lập bản đồ dữ liệu chữ theo hàng (Cột D là cột số 4 - Căn cước)
    cccd_map = {}
    data_dict = {}
    for row in range(2, sheet.max_row + 1):
        cell_val = sheet.cell(row=row, column=4).value
        if cell_val is not None:
            cccd_key = str(cell_val).strip()
            cccd_map[row] = cccd_key
            data_dict[cccd_key] = {
                "hoTen": str(sheet.cell(row=row, column=2).value or '').strip(),
                "ngaySinh": str(sheet.cell(row=row, column=3).value or '').strip(),
                "cccd": cccd_key,
                "hangGPLX": str(sheet.cell(row=row, column=5).value or '').strip(),
                "soTien": str(sheet.cell(row=row, column=6).value or '').strip()
            }

    print("-> Dang su dung openpyxl doc truc tiep anh tu o luoi...")
    success_count = 0

    # 3. Quét chính xác cấu trúc ảnh chèn trong lưới của openpyxl
    if hasattr(sheet, '_images') and len(sheet._images) > 0:
        for img in sheet._images:
            row_num = None
            if img.anchor:
                if hasattr(img.anchor, 'row'):
                    row_num = img.anchor.row + 1
                elif hasattr(img.anchor, 'from') and hasattr(img.anchor.from_, 'row'):
                    row_num = img.anchor.from_.row + 1
            
            # Mẹo đồng bộ thứ tự hàng nếu neo ô bị lỏng
            if row_num is None or row_num <= 1:
                available_rows = sorted(list(cccd_map.keys()))
                if success_count < len(available_rows):
                    row_num = available_rows[success_count]

            if row_num in cccd_map:
                cccd_name = cccd_map[row_num]
                output_path = os.path.join(output_dir, f"{cccd_name}.png")
                
                try:
                    # Đọc luồng byte ảnh thô trực tiếp
                    image_data = img.ref if hasattr(img, 'ref') else img.image
                    if isinstance(image_data, Image.Image):
                        image_data.save(output_path, "PNG")
                    else:
                        image_data.seek(0)
                        raw_bytes = image_data.read()
                        img_obj = Image.open(io.BytesIO(raw_bytes))
                        img_obj.save(output_path, "PNG")
                    success_count += 1
                except Exception as e:
                    print(f"Loi ghi anh tai hang {row_num}: {e}")

    # 4. Xuất file data.json
    with open(output_json, 'w', encoding='utf-8') as f:
        json.dump(data_dict, f, ensure_ascii=False, indent=2)
        
    print(f"\n==> HOAN THANH DONG BO! Da bock va ghi thanh cong {success_count} anh QR dang PNG.")
